"""学生成绩核心接口及考核管理页扩展接口。"""
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Literal

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Path,
    Query,
    Response,
    UploadFile,
)
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from api.common import ERROR_RESPONSES, page_data, round_2, success
from dao import score as score_dao
from dao import students_dao
from database import get_db
from scheme.common import ApiResponse, PageData
from scheme.score import (
    ScoreCreate,
    ScoreImportResult,
    ScoreOverview,
    ScoreRead,
    ScoreSummary,
    ScoreUpdate,
)


# [修改说明] 原代码：核心成绩接口返回裸模型，列表不分页，PATCH 可修改 exam_no，并另设不兼容的 /scores/search；修改原因：与统一响应及“学生+序次是业务定位键”不一致，且前端全员大表、汇总和 Excel 工具均不可用；修改为：保留 api/score.py，核心五接口对齐 backend，并在同一路由补齐前端成绩管理扩展。
router = APIRouter(tags=["学生成绩"], responses=ERROR_RESPONSES)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
MAX_EXPORT_ROWS = 50000


@dataclass(frozen=True)
class ScoreFilterState:
    filters: dict
    sort: str


def _parse_keys(value: str | None) -> list[tuple[int, int]] | None:
    if value is None:
        return None
    result = []
    seen = set()
    for raw_key in value.split(","):
        parts = raw_key.split(":")
        if len(parts) != 2:
            continue
        try:
            student_id, exam_no = (int(parts[0]), int(parts[1]))
        except ValueError:
            continue
        key = (student_id, exam_no)
        if student_id > 0 and exam_no > 0 and key not in seen:
            seen.add(key)
            result.append(key)
    return result


def get_score_filter_state(
    q: str | None = Query(default=None, max_length=100),
    class_id: int | None = Query(default=None, ge=1),
    student_id: int | None = Query(default=None, ge=1),
    exam_no: int | None = Query(default=None, ge=1),
    min_score: Decimal | None = Query(default=None, ge=0, le=100),
    max_score: Decimal | None = Query(default=None, ge=0, le=100),
    score_band: Literal["excellent", "pass", "fail"] | None = Query(
        default=None
    ),
    sort: Literal["latest", "score_desc", "score_asc", "exam_asc"] = Query(
        default="latest"
    ),
    keys: str | None = Query(default=None, max_length=10000),
) -> ScoreFilterState:
    """校验并统一全员成绩大表的筛选参数。"""
    if min_score is not None and max_score is not None and min_score > max_score:
        raise HTTPException(status_code=422, detail="最低分不能大于最高分")
    return ScoreFilterState(
        filters={
            "q": q,
            "class_id": class_id,
            "student_id": student_id,
            "exam_no": exam_no,
            "min_score": min_score,
            "max_score": max_score,
            "score_band": score_band,
            "keys": _parse_keys(keys),
        },
        sort=sort,
    )


def _grade(value) -> str:
    numeric = Decimal(value)
    if numeric >= 80:
        return "优秀"
    if numeric >= 60:
        return "及格"
    return "不及格"


def _to_overview(score, student, class_info) -> ScoreOverview:
    return ScoreOverview(
        id=score.id,
        student_id=student.id,
        student_no=student.student_no,
        name=student.name,
        class_id=class_info.id,
        class_no=class_info.class_no,
        class_name=class_info.class_name,
        exam_no=score.exam_no,
        score=score.score,
        grade=_grade(score.score),
        created_at=score.created_at,
        updated_at=score.updated_at,
    )


def _import_error(status_code: int, message: str, details=None) -> JSONResponse:
    return JSONResponse(
        status_code=status_code,
        content={
            "code": status_code * 100 + 1,
            "message": message,
            "details": details or [],
        },
    )


def get_visible_score_by_exam(db: Session, student_id: int, exam_no: int):
    """确认学生可见后，按学生 ID + 考核序次查询有效成绩。"""
    if students_dao.get_student_by_id(db, student_id) is None:
        raise HTTPException(status_code=404, detail="学生不存在或已删除")
    score = score_dao.get_active_score_by_exam(db, student_id, exam_no)
    if score is None:
        raise HTTPException(
            status_code=404,
            detail="该学生这一次考核的成绩不存在或已删除",
        )
    return score


@router.post(
    "/scores",
    response_model=ApiResponse[ScoreRead],
    status_code=201,
    summary="新增成绩",
)
def create_score(payload: ScoreCreate, db: Session = Depends(get_db)):
    """同一学生同一考核序次最多一条有效成绩。"""
    score, conflict = score_dao.create_score_guarded(db, payload.model_dump())
    if conflict == "student":
        raise HTTPException(status_code=404, detail="学生不存在或已删除")
    if conflict == "duplicate":
        raise HTTPException(
            status_code=409,
            detail="该学生这一次考核的成绩已存在，请按学生和考核序次修改",
        )
    return success(ScoreRead.model_validate(score))


@router.get(
    "/students/{student_id}/scores",
    response_model=ApiResponse[PageData[ScoreRead]],
    summary="查询学生成绩列表",
)
def list_student_scores(
    student_id: int = Path(gt=0),
    exam_no: int | None = Query(default=None, ge=1),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
):
    if students_dao.get_student_by_id(db, student_id) is None:
        raise HTTPException(status_code=404, detail="学生不存在或已删除")
    scores, total = score_dao.list_scores_by_student(
        db, student_id, page, page_size, exam_no=exam_no
    )
    items = [ScoreRead.model_validate(score) for score in scores]
    return success(page_data(items, total, page, page_size))


@router.get(
    "/students/{student_id}/scores/{exam_no}",
    response_model=ApiResponse[ScoreRead],
    summary="按学生和考核序次查询成绩",
)
def get_score_by_exam(
    student_id: int = Path(gt=0),
    exam_no: int = Path(ge=1),
    db: Session = Depends(get_db),
):
    score = get_visible_score_by_exam(db, student_id, exam_no)
    return success(ScoreRead.model_validate(score))


@router.patch(
    "/students/{student_id}/scores/{exam_no}",
    response_model=ApiResponse[ScoreRead],
    summary="按学生和考核序次修改分数",
)
def update_score_by_exam(
    payload: ScoreUpdate,
    student_id: int = Path(gt=0),
    exam_no: int = Path(ge=1),
    db: Session = Depends(get_db),
):
    score = get_visible_score_by_exam(db, student_id, exam_no)
    score = score_dao.update_score(db, score, {"score": payload.score})
    return success(ScoreRead.model_validate(score))


@router.delete(
    "/students/{student_id}/scores/{exam_no}",
    status_code=204,
    summary="按学生和考核序次逻辑删除成绩",
)
def delete_score_by_exam(
    student_id: int = Path(gt=0),
    exam_no: int = Path(ge=1),
    db: Session = Depends(get_db),
):
    score = get_visible_score_by_exam(db, student_id, exam_no)
    score_dao.delete_score(db, score)
    return Response(status_code=204)


@router.get(
    "/scores",
    response_model=ApiResponse[PageData[ScoreOverview]],
    summary="全员成绩分页大表",
)
def list_scores(
    state: ScoreFilterState = Depends(get_score_filter_state),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
):
    rows, total = score_dao.list_score_overviews(
        db,
        page,
        page_size,
        sort=state.sort,
        **state.filters,
    )
    items = [
        _to_overview(score, student, class_info)
        for score, student, class_info in rows
    ]
    return success(page_data(items, total, page, page_size))


@router.get(
    "/scores/summary",
    response_model=ApiResponse[ScoreSummary],
    summary="当前筛选条件成绩汇总",
)
def score_summary(
    state: ScoreFilterState = Depends(get_score_filter_state),
    db: Session = Depends(get_db),
):
    total, average, pass_count, excellent_count = score_dao.summarize_scores(
        db, **state.filters
    )
    result = ScoreSummary(
        total=int(total or 0),
        average_score=round_2(average) if average is not None else None,
        pass_count=int(pass_count or 0),
        excellent_count=int(excellent_count or 0),
    )
    return success(result)


@router.get(
    "/scores/exams",
    response_model=ApiResponse[list[int]],
    summary="有效成绩中的考核序次选项",
)
def list_score_exams(db: Session = Depends(get_db)):
    return success(score_dao.list_exam_numbers(db))


def _build_workbook(rows: list[ScoreOverview], template: bool) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成绩导入模板" if template else "考试成绩"
    headers = ["学号", "姓名", "班级", "考试序次", "分数"]
    sheet.append(headers)
    if not template:
        for row in rows:
            sheet.append(
                [
                    row.student_no,
                    row.name,
                    row.class_no,
                    row.exam_no,
                    float(row.score),
                ]
            )
    fill = PatternFill(fill_type="solid", fgColor="635BFF")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    widths = [24, 18, 22, 14, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E1"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _xlsx_response(content: BytesIO, filename: str) -> StreamingResponse:
    return StreamingResponse(
        content,
        media_type=XLSX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )


@router.get("/scores/template", summary="下载成绩导入模板")
def download_score_template(
    _state: ScoreFilterState = Depends(get_score_filter_state),
):
    return _xlsx_response(_build_workbook([], template=True), "score-template.xlsx")


@router.get("/scores/export", summary="导出当前筛选或选中的成绩")
def export_scores(
    state: ScoreFilterState = Depends(get_score_filter_state),
    db: Session = Depends(get_db),
):
    rows = score_dao.list_score_overviews_for_export(
        db,
        limit=MAX_EXPORT_ROWS + 1,
        sort=state.sort,
        **state.filters,
    )
    if len(rows) > MAX_EXPORT_ROWS:
        raise HTTPException(
            status_code=422,
            detail="一次最多导出 50,000 条，请缩小筛选范围",
        )
    items = [
        _to_overview(score, student, class_info)
        for score, student, class_info in rows
    ]
    return _xlsx_response(
        _build_workbook(items, template=False),
        "scores.xlsx",
    )


def _decimal_places(value: Decimal) -> int:
    normalized = value.normalize()
    return max(-normalized.as_tuple().exponent, 0)


@router.post(
    "/scores/import",
    response_model=ApiResponse[ScoreImportResult],
    status_code=201,
    summary="从 Excel 整批导入成绩",
)
async def import_scores(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        return _import_error(422, "仅支持 .xlsx 文件，请先下载模板")

    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        return _import_error(413, "文件不能超过 5 MB")
    try:
        workbook = load_workbook(BytesIO(content), data_only=False)
    except Exception:
        return _import_error(422, "文件不是有效的 Excel 工作簿")

    if not workbook.worksheets:
        return _import_error(422, "Excel 没有可导入的数据")
    sheet = workbook.worksheets[0]
    if sheet.max_row < 2:
        return _import_error(422, "Excel 没有可导入的数据")
    if sheet.max_row > MAX_IMPORT_ROWS + 1:
        return _import_error(422, "一次最多导入 1,000 行")

    aliases = {
        "student_no": {"学号", "学生编号", "student_no"},
        "exam_no": {"考试序次", "考核序次", "exam_no"},
        "score": {"分数", "成绩", "score"},
    }
    columns = {}
    for cell in sheet[1]:
        header = str(cell.value or "").strip()
        for field_name, choices in aliases.items():
            if header in choices and field_name not in columns:
                columns[field_name] = cell.column
    if set(columns) != set(aliases):
        return _import_error(422, "缺少学号、考试序次或分数列")

    details = []
    parsed_rows = []
    for row_number in range(2, sheet.max_row + 1):
        no_cell = sheet.cell(row_number, columns["student_no"])
        exam_cell = sheet.cell(row_number, columns["exam_no"])
        score_cell = sheet.cell(row_number, columns["score"])
        if all(
            cell.value is None or str(cell.value).strip() == ""
            for cell in (no_cell, exam_cell, score_cell)
        ):
            continue
        field = f"第 {row_number} 行"
        if any(
            cell.data_type == "f" for cell in (no_cell, exam_cell, score_cell)
        ):
            details.append(
                {"field": field, "reason": "学号、序次和分数不能使用公式"}
            )
            continue
        student_no = str(no_cell.value or "").strip()
        if not student_no:
            details.append({"field": field, "reason": "学号不能为空"})
            continue
        if isinstance(exam_cell.value, bool) or isinstance(score_cell.value, bool):
            details.append(
                {
                    "field": field,
                    "reason": "序次必须是正整数；分数为 0–100，最多两位小数",
                }
            )
            continue
        try:
            exam_decimal = Decimal(str(exam_cell.value).strip())
            score_value = Decimal(str(score_cell.value).strip())
        except (InvalidOperation, ValueError):
            details.append(
                {
                    "field": field,
                    "reason": "序次必须是正整数；分数为 0–100，最多两位小数",
                }
            )
            continue
        if (
            not exam_decimal.is_finite()
            or exam_decimal != exam_decimal.to_integral_value()
            or exam_decimal < 1
            or exam_decimal > 2147483647
            or not score_value.is_finite()
            or score_value < 0
            or score_value > 100
            or _decimal_places(score_value) > 2
        ):
            details.append(
                {
                    "field": field,
                    "reason": "序次必须是正整数；分数为 0–100，最多两位小数",
                }
            )
            continue
        parsed_rows.append(
            {
                "row_number": row_number,
                "student_no": student_no,
                "exam_no": int(exam_decimal),
                "score": score_value,
            }
        )

    student_map = score_dao.get_active_students_by_numbers(
        db, {row["student_no"] for row in parsed_rows}
    )
    candidates = []
    seen_keys = set()
    for row in parsed_rows:
        field = f"第 {row['row_number']} 行"
        student = student_map.get(row["student_no"])
        if student is None:
            details.append(
                {
                    "field": field,
                    "reason": f"学生 {row['student_no']} 不存在或已删除",
                }
            )
            continue
        key = (student.id, row["exam_no"])
        if key in seen_keys:
            details.append(
                {
                    "field": field,
                    "reason": "Excel 内同一学生的同一次考试重复",
                }
            )
            continue
        seen_keys.add(key)
        candidates.append(
            {
                **row,
                "student_id": student.id,
                "key": key,
            }
        )

    existing_keys = score_dao.get_active_score_keys(
        db, {row["key"] for row in candidates}
    )
    pending = []
    for row in candidates:
        if row["key"] in existing_keys:
            details.append(
                {
                    "field": f"第 {row['row_number']} 行",
                    "reason": "该次成绩已存在，请使用保存修改功能",
                }
            )
        else:
            pending.append(
                {
                    "student_id": row["student_id"],
                    "exam_no": row["exam_no"],
                    "score": row["score"],
                }
            )

    if details:
        return _import_error(
            422,
            "导入未执行，请修正以下行后重试",
            details,
        )
    if not pending:
        return _import_error(422, "没有有效数据行")

    try:
        score_dao.create_scores_batch(db, pending)
    except score_dao.ScoreBatchConflict:
        return _import_error(
            409,
            "导入期间成绩发生变化，整批已回滚，请刷新后重试",
        )

    result = ScoreImportResult(
        imported=len(pending),
        message=f"成功导入 {len(pending)} 条成绩",
    )
    return success(result)
